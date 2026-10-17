"""Build and test the whole Dataform project on your laptop, with no cloud account.

What it does:
  1. Generates the synthetic data if data/ is empty.
  2. Compiles the project with the Dataform CLI, exactly as BigQuery would run it.
  3. Loads the raw data into a local DuckDB file (data/local.duckdb).
  4. Translates each compiled BigQuery query to DuckDB with sqlglot and builds
     every model in dependency order.
  5. Runs every assertion and the unit test. Exits non-zero if any fail, so it
     works as a CI check.
  6. With --export, writes the dashboard tables to an Excel workbook for
     Google Sheets and Data Studio.

Run:
    pip install -r requirements.txt
    npm install -g @dataform/cli@3.0.70     # or let the script use npx
    python scripts/run_local.py --export

The translation step is for local testing. In production the same compiled SQL
runs unchanged in BigQuery.
"""

import argparse
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import duckdb
import sqlglot

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
DB_PATH = DATA / "local.duckdb"
sys.path.insert(0, str(ROOT / "scripts"))
from load_to_bigquery import CSV_TABLES, GA4_DATASET  # noqa: E402  (shared raw schemas)

DUCK_TYPES = {"STRING": "VARCHAR", "INT64": "BIGINT", "FLOAT64": "DOUBLE",
              "TIMESTAMP": "TIMESTAMP", "DATE": "DATE"}
GA4_COLUMNS = {
    "event_date": "VARCHAR",
    "event_timestamp": "BIGINT",
    "event_name": "VARCHAR",
    "event_params": "STRUCT(key VARCHAR, value STRUCT(string_value VARCHAR, int_value BIGINT, "
                    "float_value DOUBLE, double_value DOUBLE))[]",
    "user_pseudo_id": "VARCHAR",
    "platform": "VARCHAR",
    "collected_traffic_source": "STRUCT(manual_campaign_id VARCHAR, manual_campaign_name VARCHAR, "
                                "manual_source VARCHAR, manual_medium VARCHAR, manual_term VARCHAR, "
                                "manual_content VARCHAR, gclid VARCHAR)",
    "ecommerce": "STRUCT(purchase_revenue DOUBLE, transaction_id VARCHAR)",
}
DASHBOARD_TABLES = ["rpt_pacing", "rpt_pacing_daily", "rpt_platform_monthly",
                    "mart_campaign_performance", "qa_utm_mismatches"]


def compile_project():
    cli = ["dataform"] if shutil.which("dataform") else ["npx", "--yes", "@dataform/cli@3.0.70"]
    result = subprocess.run(cli + ["compile", "--json"], cwd=ROOT, capture_output=True, text=True)
    if result.returncode != 0:
        sys.exit(f"Dataform compile failed:\n{result.stdout}\n{result.stderr}")
    graph = json.loads(result.stdout[result.stdout.index("{"):])
    if graph.get("graphErrors", {}).get("compilationErrors"):
        sys.exit(f"Compilation errors: {graph['graphErrors']['compilationErrors']}")
    return graph


def load_raw(db):
    for (dataset, table), cols in CSV_TABLES.items():
        db.sql(f"CREATE SCHEMA IF NOT EXISTS {dataset}")
        types = ", ".join(f"'{name}': '{DUCK_TYPES[t]}'" for name, t in cols)
        db.sql(f"CREATE OR REPLACE TABLE {dataset}.{table} AS SELECT * FROM read_csv("
               f"'{DATA / dataset / f'{table}.csv'}', header = true, columns = {{{types}}})")
    db.sql(f"CREATE SCHEMA IF NOT EXISTS {GA4_DATASET}")
    cols = ", ".join(f"'{k}': '{v}'" for k, v in GA4_COLUMNS.items())
    db.sql(f"CREATE OR REPLACE TABLE {GA4_DATASET}.events_star AS SELECT * FROM read_json("
           f"'{DATA / 'ga4' / '*.jsonl'}', columns = {{{cols}}})")


def to_duckdb(sql):
    # `project.dataset.table` -> dataset.table; the GA4 wildcard events_* -> events_star
    sql = re.sub(r"`[^`.]+\.([^`.]+)\.([^`]+)`",
                 lambda m: f"{m.group(1)}.{m.group(2).replace('*', 'star')}", sql)
    return sqlglot.transpile(sql, read="bigquery", write="duckdb")[0]


def target(action):
    return f"{action['target']['schema']}.{action['target']['name']}"


def build_models(db, graph):
    declared = {target(d) for d in graph.get("declarations", [])}
    pending, built = list(graph["tables"]), set()
    for schema in {a["target"]["schema"] for a in graph["tables"] + graph.get("assertions", [])}:
        db.sql(f"CREATE SCHEMA IF NOT EXISTS {schema}")
    while pending:
        ready = [a for a in pending
                 if {f"{d['schema']}.{d['name']}" for d in a.get("dependencyTargets", [])} - declared <= built]
        if not ready:
            sys.exit(f"Unresolvable dependencies: {[target(a) for a in pending]}")
        for action in ready:
            kind = "VIEW" if action["type"] == "view" else "TABLE"
            try:
                db.sql(f"CREATE OR REPLACE {kind} {target(action)} AS {to_duckdb(action['query'])}")
            except Exception as e:  # show which model broke and why
                sys.exit(f"FAILED to build {target(action)}: {e}")
            built.add(target(action))
            pending.remove(action)
            print(f"  built  {kind.lower():5} {target(action)}")
    return len(built)


def normalize(rows):
    def cell(v):
        if v is None or isinstance(v, (str, bool)) or hasattr(v, "isoformat"):
            return v
        return round(float(v), 9)
    return sorted((tuple(cell(v) for v in r) for r in rows), key=repr)


def run_checks(db, graph):
    failures = 0
    for a in graph.get("assertions", []):
        rows = db.sql(to_duckdb(a["query"])).fetchall()
        failures += bool(rows)
        print(f"  {'PASS' if not rows else 'FAIL'}   {a['target']['name']}"
              + (f"  ({len(rows)} rows, e.g. {rows[0]})" if rows else ""))
    for t in graph.get("tests", []):
        got = normalize(db.sql(to_duckdb(t["testQuery"])).fetchall())
        want = normalize(db.sql(to_duckdb(t["expectedOutputQuery"])).fetchall())
        failures += got != want
        print(f"  {'PASS' if got == want else 'FAIL'}   unit test: {t['name']}"
              + ("" if got == want else f"\n         got  {got}\n         want {want}"))
    return failures, len(graph.get("assertions", [])) + len(graph.get("tests", []))


def export_workbook(db, path):
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name in DASHBOARD_TABLES:
            df = db.sql(f"SELECT * FROM ad_reporting.{name}").df()
            for col in df.columns:
                if col in ("date", "month", "as_of_date", "first_date", "last_date"):
                    df[col] = pd.to_datetime(df[col]).dt.date
                elif df[col].dtype == object and df[col].map(type).astype(str).str.contains("Decimal").any():
                    df[col] = df[col].astype(float)
                elif df[col].dtype.kind == "f":
                    df[col] = df[col].round(4)
            keys = [c for c in ("month", "date", "platform", "campaign_name", "campaign", "ad_group") if c in df]
            df.sort_values(keys).to_excel(writer, sheet_name=name, index=False)
        pd.DataFrame({"About this workbook": [
            "Synthetic data for a fictional client, Northwind Outdoor Co.. Every number is generated.",
            "Each tab is one table built by scripts/run_local.py from the Dataform models in this repo.",
            "To use: upload to Google Drive, open with Google Sheets, then add each tab as a Data Studio data source.",
            "Rate columns are correct only on their own row. Build rates in Data Studio as SUM()/SUM() calculated fields.",
        ]}).to_excel(writer, sheet_name="About", index=False)

    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=10, bold=c.row == 1)
        ws.freeze_panes = "A2"
        for i, header in enumerate([c.value for c in ws[1]], 1):
            letter = get_column_letter(i)
            ws.column_dimensions[letter].width = min(max(len(str(header)) + 2, 12), 110 if ws.title == "About" else 40)
            if header in ("date", "month", "as_of_date", "first_date", "last_date"):
                for c in ws[letter][1:]:
                    c.number_format = "yyyy-mm-dd"
    wb.save(path)
    print(f"\nWrote {path.relative_to(ROOT)}")


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--regenerate", action="store_true", help="rebuild the synthetic data first")
    parser.add_argument("--export", action="store_true", help="write the dashboard workbook")
    args = parser.parse_args()

    if args.regenerate or not (DATA / "ga4").exists():
        print("Generating synthetic data...")
        subprocess.run([sys.executable, str(ROOT / "scripts" / "generate_data.py")], check=True)

    print("Compiling the Dataform project...")
    graph = compile_project()

    DB_PATH.unlink(missing_ok=True)
    db = duckdb.connect(str(DB_PATH))
    print("Loading raw data into DuckDB...")
    load_raw(db)
    print("Building models...")
    n_models = build_models(db, graph)
    print("Running assertions and unit tests...")
    failures, n_checks = run_checks(db, graph)
    print(f"\n{n_models} models built, {n_checks - failures} of {n_checks} checks passed.")

    if args.export:
        export_workbook(db, DATA / "ad-reporting-dashboard-data.xlsx")
    db.close()
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
